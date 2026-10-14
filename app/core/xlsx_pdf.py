"""Excel ⇄ PDF 转换。

- Excel → PDF：调用本机 MS Excel（app.core.office），保留打印布局
- PDF → Excel：pdfplumber 逐页提取表格写入 xlsx（每页一个工作表）
  注意：只能提取"真表格"，扫描件/纯图片 PDF 无法提取。
"""

from __future__ import annotations

from pathlib import Path

import pdfplumber
from openpyxl import Workbook

from app.workers import CancelledError


def pdf_to_xlsx(pdf, out: Path,
                on_progress=None, on_log=None, is_cancelled=None) -> None:
    """PDF → Excel：逐页提取表格。"""
    on_progress = on_progress or (lambda _: None)
    on_log = on_log or (lambda _: None)
    is_cancelled = is_cancelled or (lambda: False)

    wb = Workbook()
    wb.remove(wb.active)
    table_pages, table_count = 0, 0
    with pdfplumber.open(pdf) as pdfd:
        total = len(pdfd.pages)
        for i, page in enumerate(pdfd.pages, 1):
            if is_cancelled():
                raise CancelledError("任务已取消")
            tables = page.extract_tables()
            if tables:
                table_pages += 1
                ws = wb.create_sheet(f"第{i}页")
                for t in tables:
                    for row in t:
                        ws.append(
                            ["" if c is None else str(c).replace("\n", " ").strip() for c in row]
                        )
                    ws.append([])
                    table_count += 1
            on_progress(int(i / total * 100))
    if table_pages == 0:
        wb.create_sheet("未检测到表格")
        on_log("未在该 PDF 中检测到表格（可能为扫描件），已生成空工作簿")
    wb.save(str(out))
    on_log(f"共提取 {table_count} 个表格（来自 {table_pages} 页）")
    on_progress(100)
