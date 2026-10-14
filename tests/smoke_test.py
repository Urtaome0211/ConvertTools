# -*- coding: utf-8 -*-
"""端到端冒烟测试：生成样例文件，跑通全部转换功能（含 COM 与 ffmpeg）。

用法：py tests/smoke_test.py
退出码 0 表示全部通过。
"""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

FIX = ROOT / "test_fixtures"
OUT = ROOT / "test_output"
FIX.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

FAILED = []


def ok(name: str, cond: bool, detail: str = ""):
    status = "PASS" if cond else "FAIL"
    print(f"[{status}] {name}" + (f" — {detail}" if detail else ""))
    if not cond:
        FAILED.append(name)


def on_log(msg: str):
    print("    LOG:", msg.strip()[:160])


def on_progress(p):
    if p is None or p % 25 == 0:
        print("    PROG:", p)


def build_fixtures():
    print("== 生成样例文件 ==")
    # 1) sample.pdf：3 页，第 2 页画表格
    import pymupdf
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 72), f"ConvertTools Test Page {i + 1}", fontsize=18)
        page.insert_text((72, 110), "The quick brown fox jumps over the lazy dog.", fontsize=11)
    p = doc[1]
    p.draw_rect(pymupdf.Rect(72, 160, 420, 268))  # 表格外框
    for x in (72, 188, 304):                      # 竖线
        p.draw_line((x, 160), (x, 268))
    for y in (160, 196, 232):                     # 横线
        p.draw_line((72, y), (420, y))
    for r, row in enumerate((("姓名", "部门", "工资"), ("张三", "研发", "12000"), ("李四", "测试", "9500"))):
        for c, text in enumerate(row):
            p.insert_text((84 + c * 116, 186 + r * 36), text, fontsize=11)
    doc.save(FIX / "sample.pdf")
    doc.close()

    # 2) sample.docx
    from docx import Document
    d = Document()
    d.add_heading("ConvertTools 测试文档", 0)
    d.add_paragraph("这是一段用于转换测试的正文内容。")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate((("A1", "B1", "C1"), ("A2", "B2", "C2"), ("A3", "B3", "C3"))):
        for c, text in enumerate(row):
            t.rows[r].cells[c].text = text
    d.save(FIX / "sample.docx")

    # 3) sample.xlsx
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(("项目", "数量", "单价"))
    ws.append(("苹果", 10, 3.5))
    ws.append(("香蕉", 5, 2.0))
    wb.save(FIX / "sample.xlsx")

    # 4) 图片
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (640, 400), "#4A7CF7")
    d2 = ImageDraw.Draw(img)
    d2.rectangle((80, 60, 560, 340), fill="#FFFFFF")
    d2.text((120, 180), "ConvertTools Image", fill="#4A7CF7")
    img.save(FIX / "sample.png")
    img.save(FIX / "sample.jpg")

    # 5) 视频（用项目自带 ffmpeg 生成）
    from app.core.ffmpeg import find_ffmpeg
    r = subprocess.run(
        [find_ffmpeg(), "-y", "-loglevel", "error",
         "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=15",
         "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
         "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest",
         str(FIX / "sample.mp4")],
        capture_output=True, text=True,
    )
    ok("生成 sample.mp4", r.returncode == 0, r.stderr.strip()[:120])
    print("== 样例文件就绪 ==")


def test_pdf_word():
    print("\n== PDF ⇄ Word ==")
    from app.core.pdf_word import pdf_to_docx
    out = OUT / "from_pdf.docx"
    pdf_to_docx(FIX / "sample.pdf", out)
    ok("PDF→Word", out.exists() and out.stat().st_size > 1000, f"{out.stat().st_size} B")

    from app.core.office import docx_to_pdf
    out2 = OUT / "from_docx.pdf"
    docx_to_pdf(FIX / "sample.docx", out2, on_progress=on_progress, on_log=on_log)
    ok("Word→PDF (COM)", out2.exists() and out2.stat().st_size > 500, f"{out2.stat().st_size} B")


def test_pdf_image():
    print("\n== PDF ⇄ 图片 ==")
    from app.core.pdf_image import images_to_pdf, pdf_to_images
    n = pdf_to_images(FIX / "sample.pdf", OUT / "imgs", "png", 100)
    pngs = sorted((OUT / "imgs").glob("*.png"))
    ok("PDF→图片", n == 3 and len(pngs) == 3, f"共 {n} 张")
    n2 = images_to_pdf(pngs, OUT / "from_images.pdf")
    ok("图片→PDF", n2 == 3 and (OUT / "from_images.pdf").exists())


def test_xlsx_pdf():
    print("\n== Excel ⇄ PDF ==")
    from app.core.office import xlsx_to_pdf
    out = OUT / "from_xlsx.pdf"
    xlsx_to_pdf(FIX / "sample.xlsx", out, on_progress=on_progress, on_log=on_log)
    ok("Excel→PDF (COM)", out.exists() and out.stat().st_size > 500, f"{out.stat().st_size} B")

    from app.core.xlsx_pdf import pdf_to_xlsx
    out2 = OUT / "from_pdf.xlsx"
    pdf_to_xlsx(FIX / "sample.pdf", out2)
    import openpyxl
    wb = openpyxl.load_workbook(out2)
    # 只有含表格的页才会生成工作表（样例中表格在第 2 页）
    ws = next((wb[s] for s in wb.sheetnames
               if wb[s].max_row >= 3 and wb[s].max_column >= 3), None)
    ok("PDF→Excel", out2.exists() and ws is not None,
       f"表 {ws.max_row}x{ws.max_column}" if ws else "未找到表格工作表")


def test_pdf_merge():
    print("\n== 多 PDF 合并 ==")
    from app.core.pdf_merge import merge_pdfs
    out = OUT / "merged.pdf"
    n = merge_pdfs([FIX / "sample.pdf", OUT / "from_docx.pdf"], out, on_progress=on_progress, on_log=on_log)
    import pymupdf
    pages = len(pymupdf.open(out))
    ok("合并", n == 2 and pages == 4, f"{pages} 页")


def test_video():
    print("\n== 视频 / 音视频 ==")
    from app.core.video import convert_video, extract_audio, extract_video, merge_av
    src = FIX / "sample.mp4"

    out = OUT / "video.avi"
    convert_video(src, out, "avi", on_progress=on_progress, on_log=on_log)
    ok("视频转换 mp4→avi", out.exists() and out.stat().st_size > 10000, f"{out.stat().st_size} B")

    out2 = OUT / "audio.mp3"
    extract_audio(src, out2, "mp3", on_progress=on_progress, on_log=on_log)
    ok("提取音频", out2.exists() and out2.stat().st_size > 500, f"{out2.stat().st_size} B")

    out3 = OUT / "silent.mkv"
    extract_video(src, out3, on_progress=on_progress, on_log=on_log)
    ok("提取无声视频", out3.exists() and out3.stat().st_size > 1000, f"{out3.stat().st_size} B")

    out4 = OUT / "remux.mp4"
    merge_av(src, out2, out4, on_progress=on_progress, on_log=on_log)
    ok("合并音视频", out4.exists() and out4.stat().st_size > 10000, f"{out4.stat().st_size} B")


def test_pdf_compress():
    print("\n== PDF 压缩 ==")
    from app.core.pdf_compress import compress_pdf
    out1 = OUT / "compressed_lossless.pdf"
    o1, n1 = compress_pdf(FIX / "sample.pdf", out1, mode="无损",
                          on_progress=on_progress, on_log=on_log)
    ok("无损压缩", out1.exists() and n1 > 0, f"{o1} → {n1} B")
    out2 = OUT / "compressed_lossy.pdf"
    o2, n2 = compress_pdf(FIX / "sample.pdf", out2, mode="有损", quality=60, dpi=72,
                          on_progress=on_progress, on_log=on_log)
    ok("有损压缩", out2.exists() and n2 > 0, f"{o2} → {n2} B")


def test_pdf_split():
    print("\n== PDF 拆分 ==")
    from app.core.pdf_split import split_pdf
    d1 = OUT / "split_single"
    n1 = split_pdf(FIX / "sample.pdf", d1, mode="single",
                   on_progress=on_progress, on_log=on_log)
    ok("每页一文件", n1 == 3 and len(list(d1.glob("*.pdf"))) == 3, f"{n1} 个")
    d2 = OUT / "split_ranges"
    n2 = split_pdf(FIX / "sample.pdf", d2, mode="ranges", ranges="1-2",
                   on_progress=on_progress, on_log=on_log)
    ok("自定义范围", n2 == 1 and len(list(d2.glob("*.pdf"))) == 1, f"{n2} 个")
    d3 = OUT / "split_every"
    n3 = split_pdf(FIX / "sample.pdf", d3, mode="every", every=2,
                   on_progress=on_progress, on_log=on_log)
    ok("每 N 页", n3 == 2, f"{n3} 个")


def test_pdf_security():
    print("\n== PDF 加密 / 水印 ==")
    import pymupdf
    from app.core.pdf_security import decrypt_pdf, encrypt_pdf, watermark_pdf
    enc = OUT / "encrypted.pdf"
    encrypt_pdf(FIX / "sample.pdf", enc, "test123", on_progress=on_progress, on_log=on_log)
    need = pymupdf.open(enc).needs_pass
    ok("加密", enc.exists() and need, f"needs_pass={need}")
    dec = OUT / "decrypted.pdf"
    decrypt_pdf(enc, dec, "test123", on_progress=on_progress, on_log=on_log)
    ok("解密", dec.exists() and not pymupdf.open(dec).needs_pass
       and len(pymupdf.open(dec)) == 3)
    wm = OUT / "watermarked.pdf"
    pages = watermark_pdf(FIX / "sample.pdf", wm, text="内部资料",
                          on_progress=on_progress, on_log=on_log)
    ok("文字水印", wm.exists() and pages == 3, f"{pages} 页")
    wm2 = OUT / "watermarked_img.pdf"
    pages2 = watermark_pdf(FIX / "sample.pdf", wm2, image_path=str(FIX / "sample.png"),
                           on_progress=on_progress, on_log=on_log)
    ok("图片水印", wm2.exists() and pages2 == 3, f"{pages2} 页")


def test_image():
    print("\n== 图片格式互转 ==")
    from PIL import Image
    from app.core.image_convert import convert_image
    o1, n1 = convert_image(FIX / "sample.png", OUT / "conv.webp", "webp", quality=70)
    ok("png→webp", (OUT / "conv.webp").exists() and n1 > 0, f"{o1} → {n1} B")
    o2, n2 = convert_image(FIX / "sample.jpg", OUT / "conv_small.jpg", "jpg", quality=60,
                           resize_mode="按宽度", resize_value=200)
    w = Image.open(OUT / "conv_small.jpg").width
    ok("jpg 缩放+压缩", w == 200 and 0 < n2 < o2, f"宽 {w}，{o2} → {n2} B")


def test_audio():
    print("\n== 音频格式互转 ==")
    from app.core.audio_convert import convert_audio
    src = OUT / "audio.mp3"  # 由 test_video 生成，故本测试排在 test_video 之后
    out1 = OUT / "conv.wav"
    convert_audio(src, out1, "wav", on_progress=on_progress, on_log=on_log)
    ok("mp3→wav", out1.exists() and out1.stat().st_size > 1000, f"{out1.stat().st_size} B")
    out2 = OUT / "conv.flac"
    convert_audio(src, out2, "flac", on_progress=on_progress, on_log=on_log)
    ok("mp3→flac", out2.exists() and out2.stat().st_size > 1000, f"{out2.stat().st_size} B")


def test_video_compress():
    print("\n== 视频压缩 ==")
    from app.core.video_compress import compress_video
    src = FIX / "sample.mp4"
    # 样例视频极小（约 30 KB），x264 重编码的封装开销可能让输出大于源文件，
    # 因此用两个不同 CRF 档位的输出对比来验证压缩行为。
    _, n1 = compress_video(src, OUT / "compressed_crf33.mp4", crf=33,
                           scale_mode="不缩放", on_progress=on_progress, on_log=on_log)
    _, n2 = compress_video(src, OUT / "compressed_crf20.mp4", crf=20,
                           scale_mode="不缩放", on_progress=on_progress, on_log=on_log)
    ok("画质档位与体积关系", 0 < n1 < n2, f"CRF33={n1} B < CRF20={n2} B")
    _, n3 = compress_video(src, OUT / "compressed_480.mp4", crf=28,
                           scale_mode="480p（854×480）",
                           on_progress=on_progress, on_log=on_log)
    ok("压缩+改分辨率", n3 > 0, f"{n3} B")


def test_gif():
    print("\n== 视频转 GIF ==")
    from app.core.gif_make import make_gif
    n = make_gif(FIX / "sample.mp4", OUT / "out.gif", fps=5, width=160,
                 on_progress=on_progress, on_log=on_log)
    ok("视频→GIF", (OUT / "out.gif").exists() and n > 1000, f"{n} B")


def test_gui():
    print("\n== GUI 离屏启动 ==")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    from app.main_window import MainWindow
    app = QApplication.instance() or QApplication(sys.argv)
    w = MainWindow()
    w.show()
    app.processEvents()
    ok("主窗口", w.stack.count() == 14 and w.windowTitle().startswith("ConvertTools"),
       f"{w.stack.count()} 个功能页")
    w.close()


def main():
    build_fixtures()
    for fn in (test_pdf_word, test_pdf_image, test_xlsx_pdf, test_pdf_merge,
               test_pdf_compress, test_pdf_split, test_pdf_security, test_image,
               test_video, test_audio, test_video_compress, test_gif, test_gui):
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            ok(fn.__name__, False, f"{type(exc).__name__}: {exc}")
    print("\n" + ("全部通过 ✅" if not FAILED else f"失败 {len(FAILED)} 项 ❌：{FAILED}"))
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
